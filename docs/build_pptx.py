#!/usr/bin/env python3
"""Generate the Scalable × Deep Dive sales AI strategy deck as PPTX."""
import math
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

INK = RGBColor(0x1D, 0x29, 0x33)
SOFT = RGBColor(0x55, 0x64, 0x6F)
TEAL = RGBColor(0x0E, 0x85, 0x78)
TEAL_BG = RGBColor(0xE9, 0xF4, 0xF2)
TEAL_LINE = RGBColor(0xBB, 0xDC, 0xD7)
COP = RGBColor(0xB2, 0x5A, 0x28)
COP_BG = RGBColor(0xF8, 0xEF, 0xE7)
COP_LINE = RGBColor(0xE5, 0xCD, 0xB9)
BG = RGBColor(0xFA, 0xFB, 0xFC)
LINE = RGBColor(0xDD, 0xE4, 0xE9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CHIP = RGBColor(0xEE, 0xF2, 0xF5)

FONT = "Noto Sans JP"
SW, SH = Inches(13.333), Inches(7.5)
ML = Inches(0.6)  # margin left/right

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]


def style_run(run, size, bold=False, color=INK, font=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', font)


def add_text(slide, x, y, w, h, lines, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    """lines: list of dicts {text,size,bold,color,space_before,space_after,line} or list of runs"""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ln.get('align', align)
        if 'space_before' in ln:
            p.space_before = Pt(ln['space_before'])
        if 'space_after' in ln:
            p.space_after = Pt(ln['space_after'])
        if 'line' in ln:
            p.line_spacing = ln['line']
        runs = ln['runs'] if 'runs' in ln else [ln]
        for r in runs:
            run = p.add_run()
            run.text = r['text']
            style_run(run, r.get('size', 12), r.get('bold', False), r.get('color', INK))
    return box


def add_rect(slide, x, y, w, h, fill, line_color=None, radius=0.06, shadow=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    try:
        shp.adjustments[0] = radius
    except Exception:
        pass
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line_color is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_color
        shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    return shp


def add_bg(slide):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    r.fill.solid()
    r.fill.fore_color.rgb = BG
    r.line.fill.background()
    r.shadow.inherit = False


def add_pill(slide, x, y, text, color, w=Inches(1.25)):
    pill = add_rect(slide, x, y, w, Inches(0.28), color, radius=0.5)
    tf = pill.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    style_run(run, 10, True, WHITE)
    return pill


def new_slide():
    s = prs.slides.add_slide(BLANK)
    add_bg(s)
    return s


def est_lines(text, chars_per_line):
    return max(1, math.ceil(len(text) / chars_per_line))


# ---------------------------------------------------------------- 1. title
s = new_slide()
add_text(s, ML, Inches(0.9), Inches(12), Inches(0.4),
         [dict(text="GOOGLE ADS SALES × AI STRATEGY", size=12, bold=True, color=SOFT)])
add_text(s, ML, Inches(1.35), Inches(12), Inches(1.9), [
    dict(text="提案とエンゲージメントを強化する", size=34, bold=True, line=1.2),
    dict(text="AI活用戦略", size=34, bold=True, line=1.2),
])
add_text(s, ML, Inches(3.35), Inches(11.5), Inches(0.8), [
    dict(text="ミッション(配信費用の増加 × 推奨プロダクト導入)に向けて、Googleツールのみで組む。", size=13, color=SOFT, line=1.4),
    dict(text="提案プロセス5段階を、2つのアプローチで整理する。", size=13, color=SOFT, line=1.4),
])
card_w = Inches(5.95)
for i, (tag, tagc, bgc, linec, name, desc) in enumerate([
    ("SCALABLE", TEAL, TEAL_BG, TEAL_LINE, "仕組みで、全顧客に広く",
     "GAS + Gemini API / Ads Scripts による自動化。接点の「量」とカバレッジを稼ぐ。"),
    ("DEEP DIVE", COP, COP_BG, COP_LINE, "重点顧客に、手厚く深く",
     "Gemini / Deep Research / NotebookLM を対話的に使う。提案の「質」と受注単価を上げる。"),
]):
    x = ML + (card_w + Inches(0.25)) * i
    add_rect(s, x, Inches(4.55), card_w, Inches(1.95), bgc, linec)
    add_pill(s, x + Inches(0.3), Inches(4.85), tag, tagc)
    add_text(s, x + Inches(0.3), Inches(5.28), card_w - Inches(0.6), Inches(1.1), [
        dict(text=name, size=15, bold=True, space_after=4),
        dict(text=desc, size=11.5, color=SOFT, line=1.35),
    ])

# ---------------------------------------------------------------- 2. definitions
s = new_slide()
add_text(s, ML, Inches(0.5), Inches(12), Inches(0.35),
         [dict(text="考え方", size=12, bold=True, color=SOFT)])
add_text(s, ML, Inches(0.9), Inches(12), Inches(0.55),
         [dict(text="2軸は排他ではなく、漏斗の関係", size=24, bold=True)])
add_text(s, ML, Inches(1.55), Inches(12), Inches(0.4),
         [dict(text="全顧客にScalableの網を張り、そこで見つかったシグナルを起点にDeep Diveへ切り替える。", size=12.5, color=SOFT)])

rows = [
    ("", "SCALABLE", "DEEP DIVE"),
    ("対象", "担当する全顧客(特にロングテール)", "大口・戦略顧客・増額余地の大きい顧客"),
    ("手段", "GAS + Gemini API / Ads Scripts で自動化", "Gemini / Deep Research / NotebookLM を対話的に"),
    ("成果物", "毎週自動で届くダイジェスト・下書き・アラート", "1社ごとに作り込んだ仮説・提案書・体験"),
    ("人の役割", "確認して送る/対応するだけ", "AIと壁打ちしながら質を作り込む"),
    ("効く先", "接点の量・カバレッジ", "提案の質・受注単価"),
]
tbl_shape = s.shapes.add_table(len(rows), 3, ML, Inches(2.1), Inches(12.13), Inches(3.6))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(1.5)
tbl.columns[1].width = Inches(5.32)
tbl.columns[2].width = Inches(5.31)
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        cell = tbl.cell(ri, ci)
        cell.margin_left = Inches(0.14)
        cell.margin_right = Inches(0.1)
        cell.margin_top = Inches(0.05)
        cell.margin_bottom = Inches(0.05)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        if ri == 0:
            cell.fill.fore_color.rgb = TEAL_BG if ci == 1 else (COP_BG if ci == 2 else WHITE)
        else:
            cell.fill.fore_color.rgb = WHITE if ri % 2 == 1 else BG
        p = cell.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = val
        if ri == 0:
            style_run(run, 11, True, TEAL if ci == 1 else (COP if ci == 2 else INK))
        elif ci == 0:
            style_run(run, 11, True, SOFT)
        else:
            style_run(run, 11.5, False, INK)

note = add_rect(s, ML, Inches(6.05), Inches(12.13), Inches(0.75), CHIP, LINE)
tf = note.text_frame
tf.margin_left = Inches(0.25)
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]
for txt, color, bold in [
    ("Scalableの網", TEAL, True), ((" (全顧客) → シグナル検知 (増額機会・温度感) → "), INK, False),
    ("Deep Diveで決めきる", COP, True), (" (重点顧客)", INK, False),
]:
    r = p.add_run(); r.text = txt; style_run(r, 12.5, bold, color)

# ---------------------------------------------------------------- 3. overview matrix
s = new_slide()
add_text(s, ML, Inches(0.5), Inches(12), Inches(0.35),
         [dict(text="全体マップ", size=12, bold=True, color=SOFT)])
add_text(s, ML, Inches(0.9), Inches(12), Inches(0.55),
         [dict(text="提案プロセス 5段階 × 2アプローチ", size=24, bold=True)])
mrows = [
    ("", "SCALABLE", "DEEP DIVE"),
    ("1. Planning", "ニュースダイジェスト / IS損失の自動検知", "Deep Research / NotebookLM顧客ノートブック"),
    ("2. 商談準備", "商談前ブリーフィング自動生成", "課題仮説 / 反論処理ロールプレイ"),
    ("3. 提案作成", "提案書の一括パーソナライズ", "翻訳レイヤー / 松竹梅3案 / レッドチームGem"),
    ("4. エンゲージメント", "エンゲージメント・エンジン / メール工場", "商談後60分フォロー / NotebookLM共有"),
    ("5. 効果検証", "週次サマリ / 異常検知 / 導入ギャップ分析", "QBRストーリー / So What変換"),
]
tbl_shape = s.shapes.add_table(len(mrows), 3, ML, Inches(1.75), Inches(12.13), Inches(4.9))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(2.4)
tbl.columns[1].width = Inches(4.87)
tbl.columns[2].width = Inches(4.86)
for ri, row in enumerate(mrows):
    for ci, val in enumerate(row):
        cell = tbl.cell(ri, ci)
        cell.margin_left = Inches(0.14)
        cell.margin_right = Inches(0.1)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        if ri == 0:
            cell.fill.fore_color.rgb = TEAL_BG if ci == 1 else (COP_BG if ci == 2 else WHITE)
        elif ci == 0:
            cell.fill.fore_color.rgb = CHIP
        else:
            cell.fill.fore_color.rgb = WHITE
        p = cell.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = val
        if ri == 0:
            style_run(run, 11, True, TEAL if ci == 1 else (COP if ci == 2 else INK))
        elif ci == 0:
            style_run(run, 11.5, True, INK)
        else:
            style_run(run, 11, False, INK)

# ---------------------------------------------------------------- 4-8. process slides
STEPS = ["1. Planning", "2. 商談準備", "3. 提案作成", "4. エンゲージメント", "5. 効果検証"]

procs = [
    dict(step=0, title="Planning — 誰に・なぜ・いつ攻めるか",
         scal=[("IS損失(予算)の自動検知", "Ads Scriptsで機会損失アカウントを毎週リストアップ = 増額ターゲットの自動発見", True),
               ("顧客ニュース週次ダイジェスト", "全顧客の新商品・出店・決算をGASが収集、月曜朝に「話しかけるネタ」が届く", False),
               ("伸びしろ自動スコアリング", "顧客リストを毎週採点し「今月攻める上位5社」を通知", False)],
         deep=[("NotebookLM 顧客ノートブック", "IR資料・過去提案・議事録を投入し、根拠付きでアカウントプランを構築", True),
               ("Deep Research 徹底調査", "決算・競合・市場動向を1社数十分でレポート化", False),
               ("競合出稿分析", "Ads Transparency Centerの出稿をGeminiで分析し、差別化ポイントを設計", False)],
         scal_tools="GAS + Gemini API / Ads Scripts / Sheets",
         deep_tools="Deep Research / NotebookLM / Gemini"),
    dict(step=1, title="商談準備 — 準備の速さと深さを両立する",
         scal=[("商談前ブリーフィング自動生成", "Calendarをスキャンし、全商談に「会社概要・ニュース・想定課題・推奨プロダクト」を朝に自動配信。毎回15〜30分の準備を節約", True),
               ("社内事例のNotebookLM化", "事例集を一度整備すれば、全商談で「同業種の成功事例」を引用付きで即答", False)],
         deep=[("課題仮説ジェネレーター", "顧客情報から「課題仮説3つ+根拠+検証質問」を生成。プロダクト起点でなく課題起点の商談に", True),
               ("反論処理ロールプレイ", "Geminiに厳しい決裁者を演じさせて模擬商談。「PMaxはブラックボックス」等の定番反論を事前に鍛える", False),
               ("役職別の想定問答集", "CFO/マーケ責任者/現場、それぞれの想定質問と模範回答を準備", False)],
         scal_tools="GAS + Calendar + Gemini API / NotebookLM",
         deep_tools="Gemini (Gems)"),
    dict(step=2, title="提案作成 — 「Googleの話」を「顧客のビジネスの話」に",
         scal=[("提案書の一括パーソナライズ", "マスターテンプレ+Sheets顧客データ → GASで10社分を自動差し込み生成", True),
               ("業種別プロダクト説明の量産", "「PMaxを◯◯業の言葉で」の説明バリエーションを事前に量産して使い回す", False),
               ("役職別の書き分け", "経営者3行 / マーケ半ページ / 運用者手順書、を機械的に生成", False)],
         deep=[("レッドチームGem", "「予算に厳しい役員」を演じるGemに、提出前の提案書を叩かせて弱点を潰す", True),
               ("翻訳レイヤー", "決算資料から顧客の「公式の言葉・KPI」を抽出し、提案をその語彙で書き直す", False),
               ("松竹梅の3案構成", "増額提案をYes/Noでなく「どれにするか」の議論に変える", False),
               ("Imagen / Veo クリエイティブモック", "「御社ならこんな広告が作れる」を見せる。制作リソース不足の導入障壁も解消", False)],
         scal_tools="GAS + Slides API + Gemini API",
         deep_tools="Gemini (Gems) / Imagen / Veo / Google Vids"),
    dict(step=3, title="エンゲージメント — 「連絡する理由」を絶やさない",
         scal=[("エンゲージメント・エンジン(本命)", "ニュース / パフォーマンス変化 / 季節の3シグナル → 「連絡する理由+文面下書き」を毎週自動供給", True),
               ("パーソナライズメール工場", "顧客リスト × Gemini APIで各社専用の文面をGmail下書きに一括生成。質を保って量を10倍に", False),
               ("未返信検知・休眠再活性化", "放置スレッドと45日無接点の顧客を検知し、再エンゲージ文面を下書き", False),
               ("月次インサイトレター", "業種ごとに1本、売り込まない「業界の話」を作り同業種顧客全員に展開", False)],
         deep=[("商談後60分フォロー", "Meet議事録 → お礼+決定事項+「拾い残しトピック」を抽出し、次回商談を仕込む", True),
               ("NotebookLM共有・音声概要", "提案一式をノートブックで渡し、顧客が質問しながら理解できる形に。予習用ポッドキャストも", False),
               ("Google Vids パーソナライズ動画", "多忙な決裁者向けの「3分で見られる提案」", False)],
         scal_tools="GAS + Gmail + Gemini API / Google Trends",
         deep_tools="Gemini in Meet / NotebookLM / Google Vids"),
    dict(step=4, title="効果検証 — 検証結果を次の増額・導入の弾にする",
         scal=[("週次パフォーマンスサマリ自動生成", "全アカウントの変化点+推奨アクションを自然言語で自動配信", True),
               ("異常検知アラート", "CPA/費用の急変を「顧客に説明できる文章」付きで通知 → 聞かれる前に先回り連絡", False),
               ("導入ギャップ分析", "顧客×推奨プロダクトのマトリクスから「今月の導入提案リスト」を自動生成", False),
               ("RSAアセット改善案の量産", "低評価アセットの改善案を自動生成し「来月のテスト」として提案", False)],
         deep=[("So What変換", "IS損失 → 「機会損失額」への換算など、数字を経営の言葉に翻訳した増額ストーリー", True),
               ("QBRストーリーの作り込み", "実績+目標+議事録から「達成要因 → 学び → 来期投資提案」を構築", False),
               ("検索語句クラスタリング", "部分一致+スマート自動入札を推すための実データ根拠を作る", False),
               ("悪化時の説明資料", "要因分解と挽回プランを速く出し、減額・解約を防ぐ", False)],
         scal_tools="Ads Scripts + Gemini API / Looker Studio",
         deep_tools="Gemini / Colab / Sheets"),
]

for pr in procs:
    s = new_slide()
    # stepper
    x = ML
    for i, st in enumerate(STEPS):
        w = Inches(0.45 + 0.16 * len(st))
        on = (i == pr['step'])
        chip = add_rect(s, x, Inches(0.45), w, Inches(0.32), INK if on else CHIP,
                        None if on else LINE, radius=0.5)
        tf = chip.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = Inches(0.06)
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = st
        style_run(r, 9.5, True, BG if on else SOFT)
        x += w + Inches(0.09)
    add_text(s, ML, Inches(0.95), Inches(12.1), Inches(0.55),
             [dict(text=pr['title'], size=22, bold=True)])

    col_w = Inches(5.95)
    col_top = Inches(1.7)
    col_h = Inches(5.45)
    for ci, (items, tools, tagname, tagc, bgc, linec) in enumerate([
        (pr['scal'], pr['scal_tools'], "SCALABLE", TEAL, TEAL_BG, TEAL_LINE),
        (pr['deep'], pr['deep_tools'], "DEEP DIVE", COP, COP_BG, COP_LINE),
    ]):
        cx = ML + (col_w + Inches(0.25)) * ci
        add_rect(s, cx, col_top, col_w, col_h, bgc, linec)
        add_pill(s, cx + Inches(0.25), col_top + Inches(0.22), tagname, tagc)
        iy = col_top + Inches(0.68)
        item_w = col_w - Inches(0.5)
        for name, desc, star in items:
            dlines = est_lines(desc, 36)
            ih = Inches(0.34 + 0.24 * dlines + 0.12)
            card = add_rect(s, cx + Inches(0.25), iy, item_w, ih, WHITE)
            if star:
                bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx + Inches(0.25), iy,
                                         Inches(0.05), ih)
                bar.fill.solid(); bar.fill.fore_color.rgb = tagc
                bar.line.fill.background(); bar.shadow.inherit = False
            tf = card.text_frame
            tf.word_wrap = True
            tf.margin_left = Inches(0.18)
            tf.margin_right = Inches(0.12)
            tf.margin_top = Inches(0.07)
            tf.margin_bottom = Inches(0.05)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            r = p.add_run(); r.text = name; style_run(r, 12, True, INK)
            p2 = tf.add_paragraph(); p2.line_spacing = 1.15
            p2.alignment = PP_ALIGN.LEFT
            r = p2.add_run(); r.text = desc; style_run(r, 10, False, SOFT)
            iy += ih + Inches(0.12)
        add_text(s, cx + Inches(0.25), col_top + col_h - Inches(0.4), item_w, Inches(0.3),
                 [dict(text="使う道具: " + tools, size=9.5, color=SOFT)])

# ---------------------------------------------------------------- 9. mission
s = new_slide()
add_text(s, ML, Inches(0.5), Inches(12), Inches(0.35),
         [dict(text="ミッションへの接続", size=12, bold=True, color=SOFT)])
add_text(s, ML, Inches(0.9), Inches(12), Inches(0.55),
         [dict(text="Scalableで発見し、Deep Diveで決めきる", size=24, bold=True)])

for i, (title, s_head, s_body, d_head, d_body) in enumerate([
    ("ミッション1: 配信費用の増加",
     "SCALABLE — 機会を漏らさず発見", "IS損失の自動検知 / エンゲージメント・エンジンが増額シグナルを毎週供給",
     "DEEP DIVE — 決めきる", "松竹梅3案構成 + So What変換(機会損失額)で投資判断を引き出す"),
    ("ミッション2: 推奨プロダクト導入",
     "SCALABLE — ターゲットを自動抽出", "導入ギャップ分析で「未導入×インパクト大」の顧客リストを毎月生成",
     "DEEP DIVE — 導入障壁を潰す", "翻訳レイヤーで価値を自分事化 + クリエイティブモックで制作リソース不足を解消"),
]):
    card_w = Inches(5.95)
    x = ML + (card_w + Inches(0.25)) * i
    add_rect(s, x, Inches(1.8), card_w, Inches(4.9), WHITE, LINE)
    add_text(s, x + Inches(0.35), Inches(2.1), card_w - Inches(0.7), Inches(0.4),
             [dict(text=title, size=15, bold=True)])
    stage_w = card_w - Inches(0.7)
    for j, (head, body, tagc, bgc, linec) in enumerate([
        (s_head, s_body, TEAL, TEAL_BG, TEAL_LINE),
        (d_head, d_body, COP, COP_BG, COP_LINE),
    ]):
        sy = Inches(2.65 + j * 1.95)
        stage = add_rect(s, x + Inches(0.35), sy, stage_w, Inches(1.5), bgc, linec)
        tf = stage.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.18)
        tf.margin_right = Inches(0.15)
        tf.margin_top = Inches(0.12)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run(); r.text = head; style_run(r, 10.5, True, tagc)
        p2 = tf.add_paragraph(); p2.line_spacing = 1.25; p2.space_before = Pt(3)
        p2.alignment = PP_ALIGN.LEFT
        r = p2.add_run(); r.text = body; style_run(r, 11.5, False, INK)
        if j == 0:
            add_text(s, x + Inches(0.35), sy + Inches(1.52), stage_w, Inches(0.4),
                     [dict(text="▼", size=11, color=SOFT, align=PP_ALIGN.CENTER)])

# ---------------------------------------------------------------- 10. roadmap
s = new_slide()
add_text(s, ML, Inches(0.5), Inches(12), Inches(0.35),
         [dict(text="NEXT STEPS", size=12, bold=True, color=SOFT)])
add_text(s, ML, Inches(0.9), Inches(12), Inches(0.55),
         [dict(text="実装ロードマップ", size=24, bold=True)])
add_text(s, ML, Inches(1.55), Inches(12), Inches(0.4),
         [dict(text="Deep Dive側はノーコードで今日から。コードを書く価値があるのはScalable側。", size=12.5, color=SOFT)])

phases = [
    ("今日", COP, "Gems登録(ノーコード)", "レッドチームGem / 課題仮説ジェネレーター / 翻訳レイヤー"),
    ("1週目", TEAL, "パーソナライズメール工場", "GAS + Sheets + Gemini API + Gmail下書き。小さく作れてすぐ効果が見える"),
    ("2週目", TEAL, "商談前ブリーフィング自動生成", "Calendar連動で全商談に自動配信"),
    ("3–4週目", TEAL, "週次サマリ+IS損失検知 → エンゲージメント・エンジン", "メール工場の生成部品を再利用して構築"),
]
py = Inches(2.15)
for when, wc, what, desc in phases:
    card = add_rect(s, ML, py, Inches(12.13), Inches(0.85), WHITE, LINE)
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_top = Inches(0.1)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = when + "   "; style_run(r, 13, True, wc)
    r = p.add_run(); r.text = what; style_run(r, 13, True, INK)
    p2 = tf.add_paragraph(); p2.space_before = Pt(2)
    p2.alignment = PP_ALIGN.LEFT
    r = p2.add_run(); r.text = desc; style_run(r, 10.5, False, SOFT)
    py += Inches(1.0)

add_text(s, ML, Inches(6.45), Inches(12.13), Inches(0.8), [
    dict(text="運用ルール: 自動化は「下書きまで」— 送信は必ず人間が行う / AIの数値・事例は一次ソースで裏取りしてから顧客に提示 / 顧客データの取り扱いは社内ポリシーに従う",
         size=10, color=SOFT, line=1.4),
])

out = "/tmp/claude-0/-home-user-cookie/587115f6-f8f2-5403-8cbb-b33966b7b35d/scratchpad/ai-sales-strategy.pptx"
prs.save(out)
print("saved:", out)
